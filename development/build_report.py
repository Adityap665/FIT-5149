"""Build the editable five-page report from checked data and saved experiment results.

Use the bundled document Python runtime. This does not train or change any model.
"""

from pathlib import Path
import hashlib
import json
import os

import numpy as np
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "report"
ASSETS = ROOT / "development" / "report_assets"
OUT.mkdir(exist_ok=True)
ASSETS.mkdir(exist_ok=True)
train = pd.read_csv(ROOT / "training_set.csv")
test = pd.read_csv(ROOT / "kaggle_test_X.csv")
clean = train.replace({c: {-999: np.nan} for c in
                       ["existing_repayments_aud", "has_co_applicant", "property_value_aud"]})
final = json.loads((ROOT / "development/model_checks/final_verification.json").read_text())
assert final["fresh_kernel_passed"] and final["exact_csv_match"]
assert hashlib.sha256((ROOT / "submission.csv").read_bytes()).hexdigest() == final["prediction_sha256"]
assert not train.application_id.isin(test.application_id).any()
assert train.duplicated().sum() == test.duplicated().sum() == 0
assert train.application_id.is_unique and test.application_id.is_unique
overlap = train[["property_age_years", "annual_income_aud"]].dropna()
assert len(overlap) == 11524
assert (overlap.property_age_years == overlap.annual_income_aud).all()
assert (train.sanctioned_amount_aud == 0).sum() == 5177
assert (train.sanctioned_amount_aud <= train.requested_amount_aud).all()
assert [int((train[c] == -999).sum()) for c in
        ["existing_repayments_aud", "has_co_applicant", "property_value_aud"]] == [100,107,199]


def result(name):
    return json.loads((ROOT / "development/model_checks" / name).read_text())


# A compact plot of the final model's held-out errors, not an earlier model's errors.
img = Image.new("RGB", (1600, 515), "white")
draw = ImageDraw.Draw(img)
font_path = "/System/Library/Fonts/Supplemental/Arial.ttf"
font = ImageFont.truetype(font_path, 28)
small = ImageFont.truetype(font_path, 25)
left, right, top, bottom = 120, 1510, 30, 375
for tick in [0, 20000, 40000, 60000]:
    y = bottom - tick / 65000 * (bottom - top)
    draw.line((left,y,right,y), fill="#DDDDDD", width=2)
    draw.text((left-16,y), f"{tick:,}", font=small, fill="black", anchor="rm")
for i,row in enumerate(final["error_groups"]["request_band"]):
    cx = left + (i + .5) * (right-left)/5
    y = bottom - row["rmse_aud"] / 65000 * (bottom-top)
    draw.rectangle((cx-65,y,cx+65,bottom), fill="#465C6B")
    draw.text((cx,y-10), f"{row['rmse_aud']:,.0f}", font=font, fill="black", anchor="mb")
    draw.text((cx,bottom+18), f"Q{i+1}", font=font, fill="black", anchor="mt")
draw.text((left,bottom+67), "Smallest requests", font=small, fill="black")
draw.text((right,bottom+67), "Largest requests", font=small, fill="black", anchor="ra")
draw.text((left,495), "Held-out RMSE in AUD", font=small, fill="black", anchor="ld")
chart_path = ASSETS / "final_error_by_request.png"
img.save(chart_path)

doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.28)
for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"]:
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0,0,0)
    style.font.underline = False
    # Word's built-in title style can carry a coloured rule from its template.
    for border in style.element.xpath("./w:pPr/w:pBdr"):
        border.getparent().remove(border)
normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.07
normal.paragraph_format.space_after = Pt(6)
for name,size in [("Title",20),("Heading 1",15),("Heading 2",11.5)]:
    st=doc.styles[name]
    st.font.size=Pt(size)
    st.font.bold=True
    st.paragraph_format.space_before=Pt(9 if name!="Title" else 0)
    st.paragraph_format.space_after=Pt(6)
doc.styles["Caption"].font.size=Pt(9.5)
doc.styles["Caption"].font.italic=False
doc.styles["Caption"].font.bold=False
doc.styles["Caption"].paragraph_format.space_after=Pt(6)
doc.core_properties.title="Sanctioned loan amount prediction"
doc.core_properties.author="Aditya Palve and Eric Andres Acosta Ramirez"
doc.core_properties.subject="FIT5149 Assessment 1 Group 43"

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
run=footer.add_run("Group 43 | FIT5149 Assessment 1 | ")
run.font.size=Pt(9)
field=OxmlElement("w:fldSimple");field.set(qn("w:instr"),"PAGE");footer._p.append(field)


def para(text,style=None):
    p=doc.add_paragraph(text,style)
    p.paragraph_format.widow_control=True
    return p


def heading(text,level=1):
    return doc.add_heading(text,level)


def page(title):
    doc.add_page_break()
    heading(title)


def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for cell,w in zip(t.rows[0].cells,widths): cell.width=Inches(w)
    for j,text in enumerate(headers): t.rows[0].cells[j].text=text
    for row in rows:
        cells=t.add_row().cells
        for j,value in enumerate(row): cells[j].text=str(value);cells[j].width=Inches(widths[j])
    borders=OxmlElement("w:tblBorders")
    for side in ["top","left","bottom","right","insideH","insideV"]:
        border=OxmlElement("w:"+side)
        for key,value in [("val","single"),("sz","4"),("color","D9D9D9")]: border.set(qn("w:"+key),value)
        borders.append(border)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr()
        keep=OxmlElement("w:cantSplit");trpr.append(keep)
        if i==0:
            repeat=OxmlElement("w:tblHeader");trpr.append(repeat)
        for j,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"404040" if i==0 else ("F3F4F5" if i%2==0 else "FFFFFF"));tcpr.append(shade)
            margins=OxmlElement("w:tcMar")
            for side,amount in [("top",65),("bottom",65),("left",85),("right",85)]:
                item=OxmlElement("w:"+side);item.set(qn("w:w"),str(amount));item.set(qn("w:type"),"dxa");margins.append(item)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing=1.02
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(10)
                    r.font.bold=i==0
                    r.font.color.rgb=RGBColor(255,255,255) if i==0 else RGBColor(0,0,0)
    para("").paragraph_format.space_after=Pt(0)
    return t


para("Sanctioned loan amount prediction","Title")
p=para("FIT5149 Assessment 1 | Group 43")
p.paragraph_format.space_after=Pt(3)
p=para("Aditya Palve 35324074 | Eric Andres Acosta Ramirez 32642083")
for r in p.runs:r.font.size=Pt(10)
p=para("Kaggle team TEAM ERIC-ADITYA\nUsernames adityapalve9979 and ericacostaramirez")
for r in p.runs:r.font.size=Pt(10)
heading("1 Problem and main result")
para("We predict the sanctioned amount in AUD for 10,000 unseen loan applications. Our final model is two-stage CatBoost: it estimates approval probability and the likely approved share separately. It achieved mean five-fold RMSE of AUD 33,958.77 and public Kaggle RMSE of AUD 33,096.96. We selected it using internal validation; the private Kaggle score is still unknown.")
heading("2 What we found in the data")
para("The training set has 19,322 rows and 29 columns, including the target. The test set has 28 input columns. We found no duplicate rows or IDs and no IDs shared between the sets. All labelled sanctioned amounts are between zero and the requested amount.")
para("There are 5,177 zero-sanction cases, or 26.79% of training rows. Rounded sanctioned-to-requested ratios take eight values: 0%, 65%, 70%, 75%, 80%, 85%, 90% and 100%. This pattern suggested modelling the lending decision and the approved share rather than only the dollar amount.")
rows=[]
for c in ["property_age_years","monthly_income_aud","annual_income_aud","employment_sector","credit_rating"]:
    rows.append([c,f"{train[c].isna().mean()*100:.2f}%",f"{test[c].isna().mean()*100:.2f}%"])
table(["Raw missing values","Train","Test"],rows,[4.1,1.5,1.5])
para("The missingness rates are similar across train and test, but this does not mean gaps are random or harmless. We kept missing-value flags so the model can distinguish an unavailable value from an observed one.")
para("We also found 406 occurrences of -999: 100 in repayments, 107 in co-applicant status and 199 in property value. These codes were changed to missing. Property age matched annual income exactly in all 11,524 overlapping rows and had no values from 0 to 200, so we dropped its numeric values rather than interpret them as ages.")

page("3 Cleaning and feature choices")
para("We apply the same row-by-row cleaning to training and test data. Application ID is kept only for the submission. We also exclude property_ref, an internal reference, and the corrupted property-age values. The final matrix contains 49 predictors, including 12 categorical inputs.")
para("Annual and monthly income are put on the same annual scale and combined using the median of the available values within each row. We retain the original fields, count the available income sources, and record their relative disagreement. Loan-to-income, loan-to-property-value and repayment-to-monthly-income ratios describe affordability and security; property value minus requested amount measures the remaining property value.")
para("Dates are represented by year and sine/cosine terms for month and day of year, keeping the end and start of the calendar close together. CatBoost handles numeric missing values directly. Missing categorical values receive a separate string label. In the earlier Ridge, Random Forest and HGB comparisons, imputation and encoding were fitted inside each training fold; Ridge also used scaling.")
para("We checked whether selected preprocessing decisions helped, instead of assuming that more features would improve the score. These are controlled comparisons within the stated model, not estimates of each feature's value in every model.")
table(["Controlled comparison","Mean fold RMSE in AUD"],[
 ["HGB with cleaned raw features","36,111"],
 ["HGB with ratios and missing flags","36,065"],
 ["HGB also using property age as an income backup","36,074"],
 ["Random Forest with / without property_ref","35,973 / 35,965"],
],[4.85,2.25])
para("The ratios and flags improved HGB by only AUD 46. Adding the income backup made it slightly worse. Removing property_ref changed Random Forest by just AUD 8. We kept the simpler choices, but these small differences do not establish meaningful or statistically significant gains.")
heading("4 Validation and error measures")
para("We use five shuffled StratifiedKFold splits with seed 42, balanced by the observed sanction-rate groups. Both stages train only on the fitting rows of each fold. The true sanction rate is never an input feature. Every application receives a held-out prediction, and the same folds are reused across candidates.")
para("RMSE in AUD is the selection metric because it is used by Kaggle and penalises large dollar mistakes. MAE describes the average absolute error and provides another view. MAPE would be unsuitable with 26.79% zero targets. We report mean fold RMSE; the final global OOF RMSE, calculated over all held-out predictions together, is AUD 33,975.51. Repeated model comparisons can make development scores optimistic.")

page("5 Models and final selection")
para("We began with Ridge, Random Forest and Histogram Gradient Boosting (HGB), covering a linear model, bagged trees and boosted trees. We then tested sanction-rate prediction and CatBoost. The table separates the initial direct-amount models from the later target-aware models; all use the original five folds.")
table(["Model or formulation","Mean fold RMSE in AUD"],[
 ["Ridge with prediction bounds","43,306"],
 ["Random Forest after excluding property_ref","35,965"],
 ["HGB predicting the sanction rate","34,812"],
 ["HGB using expected sanction-tier probabilities","34,625"],
 ["Single CatBoost sanction-rate regressor","34,099.76"],
 ["Two-stage CatBoost selected for attempt 4","33,958.77"],
 ["Equal average of three two-stage models","33,989.17"],
],[4.85,2.25])
para("For the selected model, a CatBoost classifier learns whether the amount is positive. A second CatBoost regressor learns the rounded sanctioned-to-requested ratio on approved fitting rows only. We multiply approval probability, predicted positive rate and requested amount. Rates and amounts are bounded to their valid ranges. Keeping probabilities avoids forcing an uncertain case into a hard approve/reject decision.")
para("The classifier uses Logloss, depth 5 and 1,200 iterations. The rate regressor uses RMSE, depth 4 and 800 iterations. Both use learning rate 0.03, L2 regularisation 10, random seed 42 and six CPU threads. The objectives follow the CatBoost documentation [1, 2].")
table(["Partition seed","Single CatBoost","Two-stage CatBoost"],[
 ["42","34,099.76","33,958.77"],
 ["2026","34,323.40","34,119.05"],
],[1.6,2.75,2.75])
para("Two-stage CatBoost improved all five paired folds in both partitions, lowering mean RMSE by AUD 141 and AUD 204 respectively. Its fold standard deviation on the original split was AUD 1,194. These are modest gains, not proof of significance. The second partition reuses the same applications and is only a stability check.")
para("A bounded comparison of 15 later configurations also tested deeper trees, more iterations, request-based weights, smaller feature sets, direct-amount CatBoost, tier classification and XGBoost. None beat the selected model on the original folds. Two fixed 50/50 blends also failed to improve global OOF RMSE. The last three-seed average was AUD 30 worse on split 42 and AUD 38 better on split 2026, effectively a tie. We retained the simpler two-stage model; no blend weights were fitted to the public leaderboard.")

page("6 Feature influence and model errors")
corr=clean[["requested_amount_aud","credit_rating","has_co_applicant"]].corrwith(clean.sanctioned_amount_aud)
para(f"We considered both relationships in the data and changes in held-out prediction error. Pearson correlations with sanctioned amount were {corr['requested_amount_aud']:.3f} for requested amount, {corr['credit_rating']:.3f} for credit rating and {corr['has_co_applicant']:.3f} for co-applicant status. These are associations, not causal effects, and dollar relationships can partly reflect the size of the loan request.")
para("For the final model, we shuffled five inputs highlighted by the EDA, one at a time in validation rows, without refitting. We repeated each shuffle three times per fold. The table reports mean RMSE increases and the standard deviation of the five fold averages; the standard deviation is not a confidence interval.")
labels={"requested_amount_aud":"Requested amount", "credit_rating":"Credit rating", "has_co_applicant":"Co-applicant status", "income_consistency":"Income consistency", "applicant_age":"Applicant age"}
table(["Input shuffled","RMSE increase AUD","Fold SD AUD"],[[labels[r['feature']],f"{r['mean']:,.0f}",f"{r['std']:,.0f}"] for r in final['permutation_checks']],[3.25,2.05,1.8])
para("Credit rating and co-applicant status matter for estimating the lending decision, while requested amount sets its dollar scale. Shuffling requested amount also changes the final conversion to AUD, so its effect describes the whole prediction process. Income consistency has a smaller effect and age very little in this model. These checks cover only five inputs; correlated features may share information, so they are not a complete ranking or a causal test.")
p=doc.add_paragraph();p.add_run().add_picture(str(chart_path),width=Inches(6.7))
p.paragraph_format.space_after=Pt(3)
para("Figure 1  Final-model OOF error by requested-amount quintile. Each group has 3,864 or 3,865 rows.","Caption")
para("Errors rise from AUD 9,014 in the smallest request group to AUD 59,394 in the largest. Rejected applications have RMSE of AUD 58,765 (n = 5,177), versus AUD 17,690 for approved applications (n = 14,145). Missing credit ratings also mark a harder group: RMSE is AUD 40,738 (n = 3,953), compared with AUD 32,006 when ratings are available (n = 15,369). These diagnostics use the final model, not the earlier Random Forest.")

page("7 Kaggle results and limitations")
table(["Attempt","Model","Public RMSE"],[
 ["1","Random Forest","35,514.14046"],
 ["2","Single CatBoost rate model","33,510.57914"],
 ["3","Same formulation in clean implementation","33,510.57914"],
 ["4","Two-stage CatBoost","33,096.96156"],
 ["5","Three-seed two-stage average","33,119.70974"],
],[.65,4.5,1.95])
para("Attempt 4 improved public RMSE by AUD 413.62 over the earlier CatBoost and was in position 19 in our screenshot. Attempt 5 was AUD 22.75 worse. Positions can change. The public score uses only 2,000 test rows; the private score on the other 8,000 determines the competition marks and was unknown when this report was written.")
para("Attempt 4's public RMSE is AUD 861.81 below its mean fold RMSE. Different samples can produce different errors, especially when a few large loans dominate RMSE. We do not treat this gap as evidence of a further improvement. The model was chosen from cross-validation before attempt 4 was uploaded. The historical attempt 1 CV log was AUD 35,956; the later saved Random Forest run is AUD 35,965, so these are not presented as one identical run.")
heading("Limits of the final model",2)
para("The final OOF MAE is AUD 16,850.52, much lower than RMSE because large mistakes receive extra weight under squared error. If we prioritised MAE, we might choose a different model. For example, the earlier hard-tier HGB had MAE around AUD 14,447 but RMSE around AUD 39,443. A lower typical absolute error did not make it better for the competition metric.")
para("The model still struggles with rejection uncertainty, large loans and missing credit ratings. The 85%, 90% and 100% tiers have only 42, 50 and 40 examples, so their behaviour is poorly supported. We would want better input data, temporal and external validation, and checks across applicant groups before considering any real use. This course dataset is synthetic; the model should not be treated as a causal measure of creditworthiness or a ready lending system.")
heading("Reproducibility and conclusion",2)
para("FIT5149_A1_final.ipynb was run from a fresh kernel and reproduced submission.csv byte for byte, with 10,000 valid rows. The run used Python 3.11.15, pandas 3.0.5, NumPy 2.4.6, scikit-learn 1.9.0 and CatBoost 1.2.10. The README records setup and paths. Earlier comparisons are retained in development files. We keep attempt 4 because it improved the original CatBoost in both validation partitions, while the final averaging experiment did not show a clear gain.")
heading("Method references",2)
para("[1] CatBoost documentation. Classification objectives, Logloss. https://catboost.ai/docs/en/concepts/loss-functions-classification\n[2] CatBoost documentation. Regression objectives, RMSE. https://catboost.ai/docs/en/concepts/loss-functions-regression","Caption")

for p in doc.paragraphs:
    assert '\u2014' not in p.text and '\u2013' not in p.text
doc.save(OUT / "group43_ass1_report.docx")
print(OUT / "group43_ass1_report.docx")
